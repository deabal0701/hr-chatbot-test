"""MUREUM 시스템 개요 - 2장을 1장으로 압축한 PPT 생성"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# 색상 정의
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_BG = RGBColor(0x1B, 0x2A, 0x4A)       # 진한 남색 배경
ACCENT_BLUE = RGBColor(0x00, 0x9E, 0xDB)    # 포인트 파란색
LIGHT_BLUE = RGBColor(0xE8, 0xF4, 0xFD)     # 연한 파란 배경
TABLE_HEADER = RGBColor(0x2C, 0x3E, 0x6B)   # 테이블 헤더
TABLE_ROW_ALT = RGBColor(0xF5, 0xF8, 0xFC)  # 테이블 줄무늬
TEXT_DARK = RGBColor(0x2D, 0x2D, 0x2D)
TEXT_GRAY = RGBColor(0x5A, 0x5A, 0x5A)
BADGE_BG = RGBColor(0xE3, 0xF0, 0xFC)
PROCESS_BLUE = RGBColor(0x00, 0x7B, 0xBE)
ARROW_COLOR = RGBColor(0xB0, 0xB0, 0xB0)


def add_section_title(slide, left, top, width, text, font_size=11):
    """섹션 제목 (파란 라운드 배지)"""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Pt(22))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT_BLUE
    shape.line.fill.background()
    shape.rotation = 0.0
    # 모서리 둥글기
    shape.adjustments[0] = 0.3
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(8)
    tf.margin_right = Pt(8)
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER


def add_table(slide, left, top, width, rows_data, col_widths, header_color=TABLE_HEADER):
    """테이블 추가 헬퍼"""
    rows = len(rows_data)
    cols = len(rows_data[0])
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, Pt(rows * 22))
    table = table_shape.table

    # 컬럼 너비 설정
    for i, w in enumerate(col_widths):
        table.columns[i].width = w

    for r_idx, row_data in enumerate(rows_data):
        for c_idx, cell_text in enumerate(row_data):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            p.text = cell_text
            p.font.size = Pt(9)
            p.font.color.rgb = WHITE if r_idx == 0 else TEXT_DARK
            p.font.bold = (r_idx == 0 or c_idx == 0)
            p.alignment = PP_ALIGN.LEFT
            cell.margin_left = Pt(6)
            cell.margin_right = Pt(4)
            cell.margin_top = Pt(3)
            cell.margin_bottom = Pt(3)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

            # 셀 배경
            if r_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = header_color
            elif r_idx % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = TABLE_ROW_ALT
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE

    return table_shape


def add_process_flow(slide, left, top, steps):
    """프로세스 흐름도 (화살표 연결)"""
    box_w = Pt(72)
    box_h = Pt(26)
    arrow_w = Pt(16)
    x = left
    for i, step in enumerate(steps):
        # 박스
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top, box_w, box_h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = PROCESS_BLUE
        shape.line.fill.background()
        shape.adjustments[0] = 0.25
        tf = shape.text_frame
        tf.margin_left = Pt(2)
        tf.margin_right = Pt(2)
        tf.margin_top = Pt(2)
        tf.margin_bottom = Pt(2)
        p = tf.paragraphs[0]
        p.text = step
        p.font.size = Pt(8)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
        x += box_w

        # 화살표 (마지막 제외)
        if i < len(steps) - 1:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW, x, top + Pt(6), arrow_w, Pt(14)
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = ARROW_COLOR
            arrow.line.fill.background()
            x += arrow_w


def add_diff_table(slide, left, top, width, rows_data, col_widths):
    """차별성 테이블 (첫 열 파란 텍스트)"""
    rows = len(rows_data)
    cols = len(rows_data[0])
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, Pt(rows * 28))
    table = table_shape.table

    for i, w in enumerate(col_widths):
        table.columns[i].width = w

    for r_idx, row_data in enumerate(rows_data):
        for c_idx, cell_text in enumerate(row_data):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            p.text = cell_text
            p.font.size = Pt(9)
            p.alignment = PP_ALIGN.LEFT
            cell.margin_left = Pt(8)
            cell.margin_right = Pt(6)
            cell.margin_top = Pt(4)
            cell.margin_bottom = Pt(4)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

            if r_idx == 0:
                # 헤더
                p.font.bold = True
                p.font.color.rgb = WHITE
                cell.fill.solid()
                cell.fill.fore_color.rgb = TABLE_HEADER
            else:
                if c_idx == 0:
                    p.font.bold = True
                    p.font.color.rgb = ACCENT_BLUE
                else:
                    p.font.color.rgb = TEXT_DARK
                cell.fill.solid()
                cell.fill.fore_color.rgb = TABLE_ROW_ALT if r_idx % 2 == 0 else WHITE

    return table_shape


def create_ppt():
    prs = Presentation()
    # 16:9 비율
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)

    # ─── 슬라이드 번호 헤더 바 ───
    header_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Pt(52)
    )
    header_bar.fill.solid()
    header_bar.fill.fore_color.rgb = DARK_BG
    header_bar.line.fill.background()

    # 슬라이드 제목
    txBox = slide.shapes.add_textbox(Pt(40), Pt(6), Inches(12), Pt(44))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    run1 = p.add_run()
    run1.text = "01  "
    run1.font.size = Pt(24)
    run1.font.bold = True
    run1.font.color.rgb = ACCENT_BLUE
    run2 = p.add_run()
    run2.text = "MUREUM 시스템 개요"
    run2.font.size = Pt(24)
    run2.font.bold = True
    run2.font.color.rgb = WHITE

    # ─── 왼쪽 영역: 시스템 개요 ───
    left_x = Pt(40)
    mid_x = Inches(6.8)

    # 설명 텍스트
    desc_box = slide.shapes.add_textbox(left_x, Pt(66), Inches(6.2), Pt(40))
    tf = desc_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = ("MUREUM은 기업 내 축적된 비정형 문서와 정형 데이터베이스를 자연어 대화만으로 "
              "통합 검색·분석·시각화할 수 있는 엔터프라이즈 AI 지식기반 어시스턴트입니다.")
    p.font.size = Pt(10)
    p.font.color.rgb = TEXT_GRAY
    p.line_spacing = Pt(16)

    # 핵심 기술 스택 배지
    add_section_title(slide, left_x, Pt(112), Pt(100), "핵심 기술 스택")

    # 기술 스택 테이블
    tech_data = [
        ["구분", "기술 내용"],
        ["AI 엔진", "LangGraph 기반 ReAct Agent + RAG + NL2SQL"],
        ["LLM", "OpenAI GPT-4o · Anthropic Claude · Google Gemini (멀티 프로바이더)"],
        ["벡터 검색", "PostgreSQL pgvector + 하이브리드 검색 (벡터 + 키워드 + RRF)"],
        ["백엔드", "FastAPI (Python 3.13) + SSE 실시간 스트리밍"],
        ["프론트엔드", "Vue 3 + ECharts 응답 시각화"],
        ["보안", "JWT 인증 + 2계층 RBAC + PII 자동 마스킹"],
    ]
    add_table(
        slide, left_x, Pt(140), Inches(6.2),
        tech_data,
        [Inches(1.0), Inches(5.2)]
    )

    # ─── 구분선 ───
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, mid_x - Pt(8), Pt(66), Pt(2), Pt(420)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(0xD0, 0xD8, 0xE8)
    line.line.fill.background()

    # ─── 오른쪽 영역: NL2SQL ───
    right_x = mid_x + Pt(8)

    # NL2SQL 제목
    nl2sql_title = slide.shapes.add_textbox(right_x, Pt(66), Inches(5.8), Pt(24))
    tf = nl2sql_title.text_frame
    p = tf.paragraphs[0]
    run1 = p.add_run()
    run1.text = "핵심 기능 — "
    run1.font.size = Pt(14)
    run1.font.bold = True
    run1.font.color.rgb = DARK_BG
    run2 = p.add_run()
    run2.text = "NL2SQL 자연어 데이터베이스 검색"
    run2.font.size = Pt(14)
    run2.font.bold = True
    run2.font.color.rgb = ACCENT_BLUE

    # 동작 프로세스 배지
    add_section_title(slide, right_x, Pt(100), Pt(90), "동작 프로세스", font_size=10)

    # 프로세스 플로우
    steps = ["사용자 질문", "의도 분석", "스키마 매칭", "SQL 생성", "검증/실행", "PII 필터", "답변 생성"]
    add_process_flow(slide, right_x, Pt(130), steps)

    # 핵심 차별성 배지
    add_section_title(slide, right_x, Pt(170), Pt(90), "핵심 차별성", font_size=10)

    # 차별성 테이블
    diff_data = [
        ["기능", "설명"],
        ["멀티턴 대화",
         '"부서별 인원수 알려줘" → "그 중 개발팀만 연봉 순으로 보여줘"처럼 이전 맥락을 이해하는 연속 대화를 지원합니다.'],
        ["의도 재작성 (Intent Rewrite)",
         "대화 이력을 분석하여 모호하거나 불완전한 질문을 완전한 질문으로 자동 재구성합니다."],
        ["자동 재시도 (Auto-Retry)",
         "SQL 오류 발생 시 오류 원인을 분석하여 자동으로 수정된 SQL을 재생성합니다."],
        ["SQL 안전장치",
         "SELECT만 허용, DDL/DML 키워드 블랙리스트, 테이블 화이트리스트, 실행 타임아웃 30초 제한."],
    ]
    add_diff_table(
        slide, right_x, Pt(198), Inches(5.8),
        diff_data,
        [Inches(1.6), Inches(4.2)]
    )

    # ─── 하단 바 ───
    footer_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, prs.slide_height - Pt(18), prs.slide_width, Pt(18)
    )
    footer_bar.fill.solid()
    footer_bar.fill.fore_color.rgb = DARK_BG
    footer_bar.line.fill.background()

    footer_text = slide.shapes.add_textbox(
        Pt(40), prs.slide_height - Pt(17), Inches(4), Pt(16)
    )
    tf = footer_text.text_frame
    p = tf.paragraphs[0]
    p.text = "MUREUM — 물어보세요, 데이터가 답합니다."
    p.font.size = Pt(8)
    p.font.color.rgb = RGBColor(0x88, 0x99, 0xBB)
    p.font.italic = True

    # 페이지 번호
    page_num = slide.shapes.add_textbox(
        prs.slide_width - Pt(60), prs.slide_height - Pt(17), Pt(40), Pt(16)
    )
    tf = page_num.text_frame
    p = tf.paragraphs[0]
    p.text = "1"
    p.font.size = Pt(8)
    p.font.color.rgb = RGBColor(0x88, 0x99, 0xBB)
    p.alignment = PP_ALIGN.RIGHT

    # 저장
    output_path = "docs/MUREUM_시스템_개요_1page.pptx"
    prs.save(output_path)
    print(f"PPT 생성 완료: {output_path}")


if __name__ == "__main__":
    create_ppt()
