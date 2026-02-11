"""Excel 내보내기 서비스

위치: app/api/services/export_service.py
- NL2SQL 조회 결과를 보고서 형태의 Excel(.xlsx)로 생성
- openpyxl 사용: 서식, 차트, 자동 열너비 지원
"""
from datetime import datetime
from io import BytesIO
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.utils.logger import setup_logger

logger = setup_logger(__name__)

# 스타일 상수
HEADER_FILL = PatternFill(start_color="2B579A", end_color="2B579A", fill_type="solid")
HEADER_FONT = Font(name="맑은 고딕", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="맑은 고딕", bold=True, size=14, color="2B579A")
SUBTITLE_FONT = Font(name="맑은 고딕", size=10, color="666666", italic=True)
DATA_FONT = Font(name="맑은 고딕", size=10)
SUMMARY_FONT = Font(name="맑은 고딕", size=9, color="888888")
THIN_BORDER = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)
ALT_ROW_FILL = PatternFill(start_color="F5F7FA", end_color="F5F7FA", fill_type="solid")

# 레이아웃 상수
MARGIN_COL = 1       # A열은 여백
START_COL = 2        # B열부터 데이터 시작
MARGIN_ROW = 1       # 1행은 여백
START_ROW = 2        # 2행부터 데이터 시작
MARGIN_WIDTH = 2     # A열(여백) 너비
TARGET_WIDTH = 130   # 전체 목표 너비 (약 1000px, 1 unit ≈ 7.7px)


class ExportService:
    """Excel 내보내기 서비스"""

    def generate_excel(
        self,
        columns: List[str],
        rows: List[Dict[str, Any]],
        question: str = "",
        sql: str = "",
        answer: str = "",
        execution_time_ms: int = 0,
        include_chart: bool = False,
        chart_config: Optional[Dict] = None,
    ) -> BytesIO:
        """보고서 형태의 Excel 파일 생성"""
        wb = Workbook()
        ws = wb.active
        ws.title = "조회결과"

        col_count = max(len(columns), 1)
        col_start = START_COL
        col_end = col_start + col_count - 1
        current_row = START_ROW

        # A열 여백
        ws.column_dimensions["A"].width = MARGIN_WIDTH

        # 열 너비 계산 (전체 약 1000px)
        available_width = TARGET_WIDTH - MARGIN_WIDTH
        col_width = max(available_width / col_count, 8)
        for ci in range(col_count):
            ws.column_dimensions[get_column_letter(col_start + ci)].width = col_width

        # 1) 제목: 질문 (첫 줄만 표시)
        if question:
            title_text = question.split("\n")[0].strip()
            ws.merge_cells(start_row=current_row, start_column=col_start, end_row=current_row, end_column=col_end)
            cell = ws.cell(row=current_row, column=col_start, value=title_text)
            cell.font = TITLE_FONT
            cell.alignment = Alignment(vertical="center")
            current_row += 1

        # 2) 부제: 생성 일시
        ws.merge_cells(start_row=current_row, start_column=col_start, end_row=current_row, end_column=col_end)
        cell = ws.cell(row=current_row, column=col_start, value=f"생성일시: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        cell.font = SUBTITLE_FONT
        cell.alignment = Alignment(horizontal="right")
        current_row += 2  # 빈 행

        # 3) 데이터 테이블 - 헤더
        header_row = current_row
        for ci, col_name in enumerate(columns):
            cell = ws.cell(row=current_row, column=col_start + ci, value=col_name)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
        current_row += 1

        # 4) 데이터 행
        numeric_columns = set()
        for row_idx, row_data in enumerate(rows):
            for ci, col_name in enumerate(columns):
                value = row_data.get(col_name, "")
                cell = ws.cell(row=current_row, column=col_start + ci, value=value)
                cell.font = DATA_FONT
                cell.border = THIN_BORDER
                if isinstance(value, (int, float)):
                    numeric_columns.add(col_start + ci)
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = '#,##0' if isinstance(value, int) else '#,##0.00'
                else:
                    cell.alignment = Alignment(horizontal="left")
                if row_idx % 2 == 1:
                    cell.fill = ALT_ROW_FILL
            current_row += 1

        data_end_row = current_row - 1
        current_row += 1  # 빈 행

        # 5) 차트 생성 (UI에서 차트 생성 버튼을 클릭한 경우만, chart_config 설정 반영)
        CHART_MAX_ROWS = 1000
        if include_chart and chart_config and len(rows) >= 2:
            chart_end_row = min(header_row + CHART_MAX_ROWS, data_end_row)
            chart_type = chart_config.get("chart_type", "line")
            x_column = chart_config.get("x_column", "")
            y_columns = chart_config.get("y_columns", [])
            pie_top_n = chart_config.get("pie_top_n", 10)

            # X축 컬럼의 Excel 열 번호
            x_col_idx = None
            if x_column in columns:
                x_col_idx = col_start + columns.index(x_column)

            # Y축 컬럼들의 Excel 열 번호
            y_col_indices = []
            for yc in (y_columns if isinstance(y_columns, list) else [y_columns]):
                if yc in columns:
                    y_col_indices.append(col_start + columns.index(yc))

            if y_col_indices:
                if chart_type == "pie":
                    # Pie 차트: 별도 데이터 시트에 Top N 데이터 작성
                    pie_col = y_col_indices[0]
                    pie_data = sorted(
                        [(str(r.get(x_column, "")), float(r.get(y_columns if isinstance(y_columns, str) else y_columns[0], 0) or 0)) for r in rows],
                        key=lambda x: x[1], reverse=True
                    )
                    if pie_top_n and pie_top_n > 0 and len(pie_data) > pie_top_n:
                        top_items = pie_data[:pie_top_n]
                        other_sum = sum(d[1] for d in pie_data[pie_top_n:])
                        top_items.append((f"기타 ({len(pie_data) - pie_top_n}건)", other_sum))
                        pie_data = top_items

                    # Pie 데이터를 임시 영역에 기록
                    pie_start_row = current_row + 1
                    for pi, (label, val) in enumerate(pie_data):
                        ws.cell(row=pie_start_row + pi, column=col_start, value=label)
                        ws.cell(row=pie_start_row + pi, column=col_start + 1, value=val)
                    pie_end_row = pie_start_row + len(pie_data) - 1

                    chart = PieChart()
                    chart.style = 10
                    chart.title = columns[pie_col - col_start]
                    chart.width = int(available_width * 0.2)
                    chart.height = 14
                    data_ref = Reference(ws, min_col=col_start + 1, min_row=pie_start_row, max_row=pie_end_row)
                    cats = Reference(ws, min_col=col_start, min_row=pie_start_row, max_row=pie_end_row)
                    chart.add_data(data_ref)
                    chart.set_categories(cats)

                    ws.add_chart(chart, f"{get_column_letter(col_start)}{current_row}")
                    # Pie 임시 데이터 + 차트 영역을 모두 건너뜀
                    hide_font = Font(color="FFFFFF", size=1)
                    for pi in range(len(pie_data)):
                        ws.cell(row=pie_start_row + pi, column=col_start).font = hide_font
                        ws.cell(row=pie_start_row + pi, column=col_start + 1).font = hide_font
                    current_row = max(current_row + 18, pie_end_row + 3)
                else:
                    # Bar / Line 차트 (여러 Y축은 하나의 차트에 시리즈로 통합)
                    chart = LineChart() if chart_type == "line" else BarChart()
                    if chart_type == "bar":
                        chart.type = "col"
                    chart.style = 10
                    chart_title = ", ".join(columns[yc - col_start] for yc in y_col_indices)
                    if len(rows) > CHART_MAX_ROWS:
                        chart_title += f" (상위 {CHART_MAX_ROWS}건)"
                    chart.title = chart_title
                    chart.width = int(available_width * 0.2)
                    chart.height = 14

                    for y_col in y_col_indices:
                        data_ref = Reference(ws, min_col=y_col, min_row=header_row, max_row=chart_end_row)
                        chart.add_data(data_ref, titles_from_data=True)

                    if x_col_idx:
                        cats = Reference(ws, min_col=x_col_idx, min_row=header_row + 1, max_row=chart_end_row)
                        chart.set_categories(cats)
                        chart.x_axis.title = columns[x_col_idx - col_start]

                    ws.add_chart(chart, f"{get_column_letter(col_start)}{current_row}")
                    current_row += 18

            current_row += 2  # 차트 후 여백

        # 6) 하단 요약 (SQL 제외)
        summary_lines = [f"조회 건수: {len(rows)}건"]
        if execution_time_ms:
            summary_lines.append(f"실행 시간: {execution_time_ms}ms")

        for line in summary_lines:
            ws.merge_cells(start_row=current_row, start_column=col_start, end_row=current_row, end_column=col_end)
            cell = ws.cell(row=current_row, column=col_start, value=line)
            cell.font = SUMMARY_FONT
            current_row += 1

        # 파일 저장
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        logger.info(f"Excel 파일 생성 완료 | rows={len(rows)}, columns={len(columns)}")
        return output


export_service = ExportService()
